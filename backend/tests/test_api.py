import pytest
from fastapi.testclient import TestClient
from main import app, seed_initial_data
from app.core.database import Base, engine, SessionLocal
from app.models.models import ConceptState, TrainingRecord, Attempt


@pytest.fixture(scope="module")
def client():
    Base.metadata.create_all(bind=engine)
    seed_initial_data()
    with TestClient(app) as test_client:
        yield test_client


def test_health_endpoint(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_get_concepts(client):
    response = client.get("/api/concepts")
    assert response.status_code == 200
    concepts = response.json()
    assert len(concepts) >= 3
    first = concepts[0]
    assert "name" in first
    assert "h_hat" in first
    assert "p_hat_current" in first


def test_concept_curve(client):
    response = client.get("/api/concepts/1/curve")
    assert response.status_code == 200
    data = response.json()
    assert data["concept_id"] == 1
    assert "curve_points" in data
    assert len(data["curve_points"]) > 10
    # Day 0 should start high
    assert data["curve_points"][0]["probability"] >= 0.95


def test_quiz_session_lifecycle(client):
    """Complete integration test:
    1. Start quiz session
    2. Submit answers (with idempotency verification)
    3. Complete session and verify HLR state updates & training records
    """
    # 1. Start quiz
    start_res = client.post("/api/quizzes/start", json={"user_id": 1, "limit": 3})
    assert start_res.status_code == 200
    quiz_data = start_res.json()
    session_id = quiz_data["session_id"]
    questions = quiz_data["questions"]
    assert len(questions) > 0

    first_q = questions[0]
    event_id = f"test-event-{session_id}-{first_q['question_id']}"

    # 2. Submit answer
    answer_res = client.post(
        f"/api/sessions/{session_id}/answers",
        json={
            "event_id": event_id,
            "question_id": first_q["question_id"],
            "selected_index": 0,
        },
    )
    assert answer_res.status_code == 200
    ans_data = answer_res.json()
    assert ans_data["already_recorded"] is False

    # 3. Test idempotency: re-submit same event_id
    duplicate_res = client.post(
        f"/api/sessions/{session_id}/answers",
        json={
            "event_id": event_id,
            "question_id": first_q["question_id"],
            "selected_index": 0,
        },
    )
    assert duplicate_res.status_code == 200
    dup_data = duplicate_res.json()
    assert dup_data["already_recorded"] is True

    # 4. Complete session
    complete_res = client.post(f"/api/sessions/{session_id}/complete")
    assert complete_res.status_code == 200
    comp_data = complete_res.json()
    assert comp_data["session_id"] == session_id
    assert comp_data["total_attempts"] >= 1
    assert len(comp_data["concepts_updated"]) >= 1

    # Verify that a training record was saved
    db = SessionLocal()
    try:
        tr = db.query(TrainingRecord).filter(TrainingRecord.session_id == session_id).first()
        assert tr is not None
        assert tr.p >= 0.0
        assert tr.x_before is not None
    finally:
        db.close()


def test_dashboard_endpoint(client):
    response = client.get("/api/dashboard?user_id=1")
    assert response.status_code == 200
    data = response.json()
    assert "overall_retention" in data
    assert "subjects" in data
    assert len(data["subjects"]) >= 1


def test_study_plan_endpoints(client):
    today_res = client.get("/api/plan/today?user_id=1")
    assert today_res.status_code == 200
    today_data = today_res.json()
    assert len(today_data["plan"]) == 1
    assert "concepts" in today_data["plan"][0]

    week_res = client.get("/api/plan/week?user_id=1")
    assert week_res.status_code == 200
    week_data = week_res.json()
    assert len(week_data["plan"]) == 7


def test_materials_upload(client):
    notes = """
    Cache Coherence: Protocols to ensure consistency among multiple local caches.
    Write-Through: Data is written to cache and main memory simultaneously.
    Write-Back: Data is written only to cache; memory is updated only on replacement.
    """
    response = client.post(
        "/api/materials",
        json={
            "subject_id": 1,
            "title": "Cache Architecture",
            "content": notes,
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert data["concepts_generated"] >= 1
    assert data["questions_generated"] >= 1


def test_settings_endpoints(client):
    get_res = client.get("/api/settings?user_id=1")
    assert get_res.status_code == 200
    assert "daily_minutes" in get_res.json()

    put_res = client.put(
        "/api/settings?user_id=1",
        json={
            "daily_minutes": 90,
            "subject_priorities": {"1": 1.5},
            "exam_dates": {"1": "2026-12-01"},
        },
    )
    assert put_res.status_code == 200
    assert put_res.json()["daily_minutes"] == 90


def test_ai_status_endpoint(client):
    res = client.get("/api/ai/status")
    assert res.status_code == 200
    data = res.json()
    assert "status" in data


def test_ai_chat_endpoint(client):
    res = client.post(
        "/api/ai/chat",
        json={
            "message": "Why does memory decay exponentially?",
            "history": [],
            "subject": "Cognitive Science",
            "persona": "socratic",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "reply" in data
    assert len(data["reply"]) > 10
    assert "suggested_followups" in data


def test_ai_doubt_solver_endpoint(client):
    res = client.post(
        "/api/ai/solve-doubt",
        json={
            "doubt": "What is the difference between Write-Through and Write-Back caching?",
            "subject": "Operating Systems",
            "difficulty": "Medium",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "core_intuition" in data
    assert len(data["step_by_step_solution"]) >= 1
    assert "retention_hook" in data
    assert "verification_challenge" in data


def test_ai_mnemonics_endpoint(client):
    res = client.post(
        "/api/ai/mnemonics",
        json={
            "concept_name": "Raft Leader Election",
            "topic": "Consensus",
            "subject": "Distributed Systems",
        },
    )
    assert res.status_code == 200
    data = res.json()
    assert "story_mnemonic" in data
    assert "acronym_or_rhyme" in data
    assert "flashcards" in data
    assert len(data["flashcards"]) >= 1


def test_extract_file_docx(client):
    import io
    import docx
    
    doc = docx.Document()
    doc.add_heading("Operating Systems Virtual Memory", level=1)
    doc.add_paragraph("Paging and segmentation are the core memory virtualization techniques.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    
    files = {"file": ("virtual_memory.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/materials/extract-file", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "virtual_memory.docx"
    assert "Paging and segmentation" in data["text"]
    assert data["file_type"] == "docx"


def test_extract_file_pptx(client):
    import io
    import pptx
    
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Distributed Systems Slide"
    slide.placeholders[1].text = "CAP Theorem dictates trade-offs between consistency and availability."
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    
    files = {"file": ("dist_systems.pptx", buf, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = client.post("/api/materials/extract-file", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "dist_systems.pptx"
    assert "CAP Theorem" in data["text"]
    assert data["file_type"] == "pptx"


def test_extract_file_txt(client):
    import io
    buf = io.BytesIO(b"Deadlock avoidance requires Banker's algorithm and resource ordering.")
    files = {"file": ("deadlocks.txt", buf, "text/plain")}
    res = client.post("/api/materials/extract-file", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["filename"] == "deadlocks.txt"
    assert "Banker's algorithm" in data["text"]
    assert data["file_type"] == "text"


def test_materials_subject_filter(client):
    res = client.get("/api/materials?subject=Operating%20Systems")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    for mat in data:
        assert mat["subject"].lower() == "operating systems"


def test_plan_tiles_exact_timestamps_and_topics(client):
    res = client.get("/api/plan/today")
    assert res.status_code == 200
    data = res.json()
    assert "tiles" in data
    assert len(data["tiles"]) >= 1

    first_tile = data["tiles"][0]
    # Verify exact timestamps on each tile
    assert "time_slot" in first_tile
    assert " - " in first_tile["time_slot"]
    assert "AM" in first_tile["time_slot"] or "PM" in first_tile["time_slot"]
    assert "start_time" in first_tile
    assert "end_time" in first_tile
    
    # Verify subject and weak topic mapping
    assert "subject" in first_tile
    assert "recommended_topic" in first_tile
    assert first_tile["recommended_topic"].startswith("Recommended Topic: ")
    assert "productivity_tier" in first_tile
    assert "mode_tier" in first_tile
    assert first_tile["mode_tier"] == "AI-Optimized"


def test_plan_schedule_free_time_windows_slicing(client):
    # Free-time window: 07:00 to 08:30 (90 mins). User selects session duration: 30 mins (Manual mode).
    # Expected: Exactly 3 uniform sliced session tiles, no micro-intervals.
    payload = {
        "mode": "manual",
        "session_duration": 30,
        "free_time_windows": [
            {"id": "win-morning", "start": "07:00", "end": "08:30", "label": "Morning Study"}
        ]
    }
    res = client.post("/api/plan/schedule", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["session_duration"] == 30
    assert len(data["tiles"]) == 3
    assert data["tiles"][0]["time_slot"] == "07:00 AM - 07:30 AM"
    assert data["tiles"][1]["time_slot"] == "07:30 AM - 08:00 AM"
    assert data["tiles"][2]["time_slot"] == "08:00 AM - 08:30 AM"
    for tile in data["tiles"]:
        assert tile["duration_minutes"] == 30
        assert tile["recommended_topic"].startswith("Recommended Topic: ")
        assert "🧠 Cognitive Peak" in tile["productivity_tier"]


def test_plan_schedule_ai_session_length_optimization(client):
    # AI Mode for session length: User provides initial duration anchor (e.g. 45m).
    # System analyzes quiz data and tunes towards optimal duration for retention (20m due to decay).
    payload = {
        "mode": "ai_optimized",
        "session_length_mode": "ai_optimized",
        "initial_session_duration": 45,
        "free_time_windows": [
            {"id": "win-morning", "start": "07:00", "end": "08:00", "label": "Morning Focus"}
        ]
    }
    res = client.post("/api/plan/schedule", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["session_length_mode"] == "ai_optimized"
    assert data["initial_session_duration"] == 45
    assert data["optimal_session_duration"] in (20, 30, 45)
    assert "ai_duration_optimization" in data
    assert data["ai_duration_optimization"]["mode"] == "ai_optimized"
    assert "reason" in data["ai_duration_optimization"]
    # Verify formatted time slot
    assert "time_slot_formatted" in data["tiles"][0]
    assert "to" in data["tiles"][0]["time_slot_formatted"]


def test_plan_schedule_manual_mode(client):
    # Manual mode with user-dictated slots
    payload = {
        "mode": "manual",
        "session_duration": 45,
        "custom_slots": [
            {
                "start_time": "09:00 PM",
                "end_time": "09:45 PM",
                "subject": "Distributed Systems",
                "target_topic": "Raft Leader Election",
                "productivity_tier": "🌙 Night Review"
            }
        ]
    }
    res = client.post("/api/plan/schedule", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["mode"] == "manual"
    assert len(data["tiles"]) == 1
    tile = data["tiles"][0]
    assert tile["time_slot"] == "09:00 PM - 09:45 PM"
    assert tile["subject"] == "Distributed Systems"
    assert "Raft Leader Election" in tile["recommended_topic"]
    assert tile["mode_tier"] == "Manual"



